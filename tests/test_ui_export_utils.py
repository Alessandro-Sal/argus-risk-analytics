"""
==============================================================================
ARGUS FINANCIAL EXPORT FRAMEWORK UNIT TESTS (tests/test_ui_export_utils.py)
==============================================================================
Test suite per la validazione delle funzionalità di esportazione dati tabulari:
- Sanitizzazione DataFrame (MultiIndex, indici complessi, valori non primitivi).
- Integrità CSV UTF-8 con BOM (utf-8-sig) per compatibilità Excel europeo.
- Integrità e styling del file Excel .xlsx (openpyxl / xlsxwriter).
- Generazione dinamica dei nomi file e caching dei buffer.
- Esecuzione sicura dei componenti UI Streamlit (render_export_toolbar, render_table_with_export).
==============================================================================
"""

import datetime
import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

from core.ui_export_utils import (
    generate_export_filename,
    get_cached_csv_bytes,
    get_cached_excel_bytes,
    prepare_dataframe_for_export,
    render_export_toolbar,
    render_table_with_export,
    to_csv_bytes,
    to_excel_bytes,
)
from core.ui_utils import render_data_table

# ==============================================================================
# 1. TEST PREPARAZIONE E SANITIZZAZIONE DATAFRAME
# ==============================================================================

def test_prepare_dataframe_none_and_empty():
    assert prepare_dataframe_for_export(None).empty
    assert prepare_dataframe_for_export(pd.DataFrame()).empty


def test_prepare_dataframe_multiindex_columns():
    tuples = [('Portafoglio', 'Ticker'), ('Performance', 'Rendimento_pct'), ('Rischio', 'VaR_95')]
    multi_cols = pd.MultiIndex.from_tuples(tuples)
    df = pd.DataFrame([["AAPL", 12.5, -3.2]], columns=multi_cols)
    
    clean = prepare_dataframe_for_export(df)
    assert not isinstance(clean.columns, pd.MultiIndex)
    assert "Portafoglio_Ticker" in clean.columns
    assert "Performance_Rendimento_pct" in clean.columns
    assert "Rischio_VaR_95" in clean.columns


def test_prepare_dataframe_named_index():
    df = pd.DataFrame({"Prezzo": [150.0, 155.0]}, index=pd.Index(["2026-01-01", "2026-01-02"], name="Data"))
    clean = prepare_dataframe_for_export(df)
    assert "Data" in clean.columns
    assert "Prezzo" in clean.columns


def test_prepare_dataframe_complex_nested_types():
    df = pd.DataFrame({
        "Asset": ["BTC", "ETH"],
        "Metadata": [{"network": "mainnet", "fee": 1.5}, {"network": "pos", "fee": 0.8}],
        "Tags": [["crypto", "core"], ["smart_contracts"]]
    })
    clean = prepare_dataframe_for_export(df)
    assert isinstance(clean["Metadata"].iloc[0], str)
    assert isinstance(clean["Tags"].iloc[0], str)
    assert "mainnet" in clean["Metadata"].iloc[0]


# ==============================================================================
# 2. TEST GENERATORE NOMI FILE DINAMICO
# ==============================================================================

def test_generate_export_filename():
    fn = generate_export_filename(prefix="posizioni/dettaglio*", portfolio_name="Global Tech & ESG", extension="xlsx")
    assert fn.endswith(".xlsx")
    assert "posizioni_dettaglio" in fn
    assert "Global_Tech___ESG" in fn
    # Verifica presenza del timestamp YYYYMMDD
    today_str = datetime.datetime.now().strftime("%Y%m%d")
    assert today_str in fn
    # Nessun carattere proibito nei filesystem
    for bad_char in r'/\:*?"<>|':
        assert bad_char not in fn


# ==============================================================================
# 3. TEST SERIALIZZATORE CSV (UTF-8 SIG CON BOM)
# ==============================================================================

def test_to_csv_bytes_bom_and_accents():
    df = pd.DataFrame({
        "Ticker": ["ENEL.MI", "ISP.MI"],
        "Descrizione": ["Società Elettrica Nazionale", "Intesa Sanpaolo S.p.A."],
        "Prezzo Carico (€)": [6.45, 2.80],
        "PnL Latente (€)": [1250.50, -320.10]
    })
    
    csv_bytes = to_csv_bytes(df, sep=",")
    assert isinstance(csv_bytes, bytes)
    assert len(csv_bytes) > 0
    
    # Verifica presenza del BOM UTF-8 (\xef\xbb\xbf)
    assert csv_bytes.startswith(b'\xef\xbb\xbf')
    
    # Decodifica con utf-8-sig e verifica integrità del simbolo dell'Euro e lettere accentate
    decoded_text = csv_bytes.decode('utf-8-sig')
    assert "Società Elettrica" in decoded_text
    assert "Prezzo Carico (€)" in decoded_text
    assert "1250.5" in decoded_text


# ==============================================================================
# 4. TEST SERIALIZZATORE EXCEL PROFESSIONALE (.XLSX)
# ==============================================================================

def test_to_excel_bytes_structure_and_styling():
    df = pd.DataFrame({
        "Ticker": ["AAPL", "MSFT", "NVDA"],
        "Controvalore (€)": [15200.50, 22400.00, 8900.75],
        "Peso (%)": [32.5, 47.9, 19.6],
        "Quantità": [80, 55, 75],
        "Data Acquisto": [datetime.date(2025, 6, 15), datetime.date(2025, 9, 20), datetime.date(2026, 1, 10)]
    })

    excel_bytes = to_excel_bytes(df, sheet_name="Posizioni_Core", base_currency="EUR")
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000

    # Riapertura workbook in-memory con openpyxl per validare stili e formati
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Posizioni_Core"

    # 1. Verifica Griglia Contabile Attiva
    assert ws.views.sheetView[0].showGridLines is True

    # 2. Verifica Freeze Panes (blocco intestazione riga 1)
    assert ws.freeze_panes == "A2"

    # 3. Verifica Header (riga 1)
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.value == "Ticker"
    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb in ("00FFFFFF", "FFFFFFFF", "FFFFFF")
    # Colore di sfondo #161B22
    assert "161B22" in str(header_cell.fill.start_color.rgb)

    # 4. Verifica Formattazione Numerica Automatica
    # Colonna 2: Controvalore (€) -> Valuta
    val_cell = ws.cell(row=2, column=2)
    assert isinstance(val_cell.value, (int, float))
    assert "€" in val_cell.number_format or "#,##0.00" in val_cell.number_format

    # Colonna 3: Peso (%) -> Percentuale
    pct_cell = ws.cell(row=2, column=3)
    assert "%" in pct_cell.number_format

    # Colonna 4: Quantità -> Intero
    qty_cell = ws.cell(row=2, column=4)
    assert qty_cell.number_format == "#,##0"

    # 5. Verifica Auto-fit larghezza colonna
    col_a_width = ws.column_dimensions["A"].width
    col_b_width = ws.column_dimensions["B"].width
    assert col_a_width >= 12
    assert col_b_width >= 12


# ==============================================================================
# 5. TEST PERFORMANCE E CACHING
# ==============================================================================

def test_caching_functions():
    df = pd.DataFrame({"X": [1, 2, 3], "Y": [10.5, 20.5, 30.5]})
    b_csv_1 = get_cached_csv_bytes(df, sep=",")
    b_csv_2 = get_cached_csv_bytes(df, sep=",")
    assert b_csv_1 == b_csv_2

    b_xlsx_1 = get_cached_excel_bytes(df, sheet_name="Test")
    b_xlsx_2 = get_cached_excel_bytes(df, sheet_name="Test")
    assert len(b_xlsx_1) > 0
    assert len(b_xlsx_2) > 0


# ==============================================================================
# 6. TEST COMPONENTI STREAMLIT UI SENZA ERRORI A RUNTIME
# ==============================================================================

def test_render_export_toolbar_and_table_with_export():
    df = pd.DataFrame({
        "Ticker": ["ISP.MI", "ENI.MI"],
        "Valore (€)": [12000.0, 8500.0],
        "Allocazione (%)": [58.5, 41.5]
    })

    # Verifica invocabilità di render_export_toolbar
    render_export_toolbar(
        df=df,
        file_prefix="test_pos",
        key_suffix="t1",
        table_title="Posizioni Test",
        show_row_count=True
    )

    # Verifica invocabilità di render_table_with_export
    render_table_with_export(
        df=df,
        table_title="Tabella Dati Istituzionale",
        file_prefix="test_tbl",
        key_suffix="t2",
        currency_cols=["Valore (€)"],
        pct_cols=["Allocazione (%)"],
        height=300
    )

    # Verifica retrocompatibilità di render_data_table estesa
    render_data_table(
        df=df,
        currency_cols=["Valore (€)"],
        download_filename="export_integrato_test",
        table_title="Tabella Integrata"
    )
